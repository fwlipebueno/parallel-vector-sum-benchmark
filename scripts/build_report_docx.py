#!/usr/bin/env python3
import argparse
import csv
import platform
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

METHOD_ORDER = ["Sequencial", "OpenMP", "TBB"]


def read_summary(path):
    with open(path, newline="", encoding="utf-8") as file:
        return list(csv.DictReader(file))


def add_heading(document, text, level):
    document.add_heading(text, level=level)


def add_paragraph(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_results_table(document, rows):
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["Abordagem", "Execuções", "Média (s)", "Desvio padrão (s)", "Mínimo (s)", "Máximo (s)", "Speedup"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header

    ordered = sorted(rows, key=lambda row: METHOD_ORDER.index(row["method"]) if row["method"] in METHOD_ORDER else 99)
    for row in ordered:
        cells = table.add_row().cells
        cells[0].text = row["method"]
        cells[1].text = row["runs"]
        cells[2].text = f"{float(row['mean_seconds']):.6f}"
        cells[3].text = f"{float(row['stdev_seconds']):.6f}"
        cells[4].text = f"{float(row['min_seconds']):.6f}"
        cells[5].text = f"{float(row['max_seconds']):.6f}"
        cells[6].text = f"{float(row['speedup_vs_sequential']):.2f}x"


def add_image_if_exists(document, path, caption):
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(5.8))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)


def best_parallel(rows):
    candidates = [row for row in rows if row["method"] in {"OpenMP", "TBB"}]
    if not candidates:
        return None
    return min(candidates, key=lambda row: float(row["mean_seconds"]))


def build(summary_path, results_dir, output_path):
    rows = read_summary(summary_path)
    best = best_parallel(rows)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_heading("Relatório técnico — soma de vetor com C++, OpenMP e TBB", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(document, "1. Introdução", 1)
    add_paragraph(document, "Este experimento compara três formas de somar os elementos de um vetor grande em C++: implementação sequencial, paralelização com OpenMP e paralelização com Intel TBB. O objetivo é observar o impacto da paralelização no tempo de execução e avaliar a estabilidade dos resultados por meio da média e do desvio padrão das medições.")

    add_heading(document, "2. Metodologia", 1)
    add_paragraph(document, "O vetor utilizado possui 100 milhões de elementos inteiros, inicializados com o valor 1. Cada abordagem foi executada 30 vezes. Em todas as execuções, a soma foi validada antes do registro do tempo.")
    add_paragraph(document, "A medição foi realizada dentro do programa C++ com std::chrono::steady_clock, registrando apenas o trecho responsável pela soma. A versão sequencial utiliza um laço for simples; a versão OpenMP utiliza parallel for com redução; e a versão TBB utiliza parallel_reduce.")
    add_paragraph(document, f"Ambiente usado para gerar este relatório: {platform.platform()}.")

    add_heading(document, "3. Resultados", 1)
    add_results_table(document, rows)

    add_heading(document, "4. Gráficos", 1)
    add_image_if_exists(document, results_dir / "comparison_times.png", "Figura 1 — Comparação dos tempos de execução.")
    add_image_if_exists(document, results_dir / "sequential_times.png", "Figura 2 — Tempos da abordagem sequencial.")
    add_image_if_exists(document, results_dir / "openmp_times.png", "Figura 3 — Tempos da abordagem OpenMP.")
    add_image_if_exists(document, results_dir / "tbb_times.png", "Figura 4 — Tempos da abordagem TBB.")

    add_heading(document, "5. Discussão", 1)
    if best:
        add_paragraph(document, f"Considerando a média das execuções, a abordagem paralela com melhor desempenho foi {best['method']}, com speedup médio de {float(best['speedup_vs_sequential']):.2f}x em relação à versão sequencial.")
    add_paragraph(document, "A soma de um vetor grande é uma operação simples e fortemente dependente da leitura de memória. Por isso, a paralelização pode reduzir o tempo de execução, mas o ganho tende a ser limitado pela largura de banda de memória e pelo custo de coordenação das threads.")
    add_paragraph(document, "OpenMP oferece uma forma direta de paralelizar laços, enquanto TBB utiliza uma abordagem baseada em divisão de intervalos e redução paralela. A diferença entre as duas bibliotecas depende da máquina, do runtime, do número de núcleos disponíveis e da carga do sistema durante o experimento.")

    add_heading(document, "6. Conclusão", 1)
    add_paragraph(document, "O experimento mostra como a paralelização pode melhorar o desempenho de uma tarefa simples em C++, mas também evidencia a importância de medir várias execuções e analisar média, desvio padrão e comportamento gráfico. A estrutura utilizada permite reproduzir o teste e comparar os resultados com maior segurança do que uma única execução isolada.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Build the DOCX report from benchmark outputs.")
    parser.add_argument("--summary", default="results/summary.csv")
    parser.add_argument("--results", default="results")
    parser.add_argument("--out", default="results/report.docx")
    args = parser.parse_args()

    build(Path(args.summary), Path(args.results), Path(args.out))
    print(f"DOCX report: {args.out}")


if __name__ == "__main__":
    main()
